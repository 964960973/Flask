import json
import time
from openpyxl import load_workbook
import threading
import jsonpath
from seleniumwire import webdriver
import requests
import redis
from lxml import etree
from openpyxl.drawing.image import Image
import os
from PIL import Image as image_change
from selenium.webdriver.common.by import By
from selenium import webdriver
import redis
from sql_help import db

redis_client = redis.Redis(host='127.0.0.1', port=6379)





def sql_table(website,hity,name):
    sql = f"""CREATE TABLE {website}_{hity}_{name} (
         id int primary key auto_increment not null,
         详情链接  varchar(200) ,
         所在地区  varchar(200),
         职位  varchar(200),
         所需技术  varchar(200),
         福利待遇  varchar(200),
         企业信息  varchar(200),
         工资信息  varchar(200))
         """
    db.tables(sql)


def nationwide(website,hity,name):
    sql = f"""CREATE TABLE {website}_{hity}_{name} (
         id int primary key auto_increment not null,
         详情链接  varchar(200) ,
         所在地区  varchar(200),
         职位  varchar(200),
         所需技术  varchar(200),
         福利待遇  varchar(200),
         企业信息  varchar(200),
         工资信息  varchar(200))
         """
    db.tables(sql)

def mobie():
    url = "https://www.zhipin.com/wapi/zpgeek/common/data/city/site.json"
    payload={}
    headers = {
        'authority': 'www.zhipin.com',
        'accept': 'application/json, text/javascript, */*; q=0.01',
        'accept-language': 'zh-CN,zh;q=0.9',
        'cache-control': 'no-cache',
        'cookie': 'lastCity=101280100; wd_guid=63ad5a73-16a5-4e1e-8f12-4189d2ef8aeb; historyState=state; _bl_uid=4wlz5dmwms9uLzkq47FXf1wdU4Ik; sid=sem_pz_bdpc_dasou_title; __zp_seo_uuid__=9ffd5a7e-1074-470b-a7cd-b0a8de5525d0; __g=sem_pz_bdpc_dasou_title; Hm_lvt_194df3105ad7148dcf2b98a91b5e727a=1680081954,1680852957,1681786419; Hm_lpvt_194df3105ad7148dcf2b98a91b5e727a=1681786423; collection_pop_window=1; __zp_stoken__=37efeJGt0ejd8DG4UbVsxFVojXX1rbD0yTXw1XzUVH3pGKjxJeHFaTSUOZhIvC1FWBWRHOF04P14vZUdkZRRjTDcqKgcnCjRtBQsTK0EUIxdHAwMDYAoAZ1B8EkJTBHgdR111QwZDTzxtVjQ%3D; __c=1681786418; __l=r=https%3A%2F%2Fwww.baidu.com%2Fother.php%3Fsc.Ks0000aqqdzCLIcx93mp_yBuV0hvZodmAUBZGGImL2_4gUSx2E7Cik9xBd3SwBwKDcKtIK9Zz8DZeVLf9Tigq56EZveER_FPm8eyg1MM_siw-4LoQA0-K2l-gqpa457Uu9SUq08cSwYdoC5_0nT_Rk7s99jc54M_hukP84AO6AroPxgPbj8q-x4eGgmQIoy6ZOno1CS9fLNS1A2VCVZRAOP6ZFJC.7D_NR2Ar5Od663rj6t8AGSPticrtXFBPrM-kt5QxIW94UhmLmry6S9wiGyAp7BEIu80.TLFWgv-b5HDkrfK1ThPGujYknHb0THY0IAYqmhq1Tqpkko60IgP-T-qYXgK-5H00mywxIZ-suHY10ZIEThfqmhq1Tqpkko60ThPv5H00IgF_gv-b5HDdPjf4PHTkrjR0UgNxpyfqnHRzn1mYnHc0UNqGujYknWDkrHRLr0KVIZK_gv-b5HDzrjcv0ZKvgv-b5H00pywW5R9rffKspyfqP0KWpyfqrjf0mLFW5HRYP1fs%26dt%3D1681786414%26wd%3Dboss%25E7%259B%25B4%25E8%2581%2598%26tpl%3Dtpl_12826_31784_0%26l%3D1544957185%26us%3DlinkVersion%253D1%2526compPath%253D10036.0-10032.0%2526label%253D%2525E4%2525B8%2525BB%2525E6%2525A0%252587%2525E9%2525A2%252598%2526linkType%253D%2526linkText%253DBOSS%2525E7%25259B%2525B4%2525E8%252581%252598%2525E2%252580%252594%2525E2%252580%252594%2525E6%252589%2525BE%2525E5%2525B7%2525A5%2525E4%2525BD%25259C%2525EF%2525BC%25258C%2525E4%2525B8%25258ABOSS%2525E7%25259B%2525B4%2525E8%252581%252598%2525EF%2525BC%25258C&l=%2Fwww.zhipin.com%2Fweb%2Fgeek%2Fjob%3Fcity%3D101280100&s=3&g=%2Fwww.zhipin.com%2Fguangzhou%2F%3Fsid%3Dsem_pz_bdpc_dasou_title&friend_source=0&s=3&friend_source=0; __a=22574609.1675327175.1680852958.1681786418.67.4.6.6',
        'pragma': 'no-cache',
        'referer': 'https://www.zhipin.com/web/geek/job?city=101280100',
        'sec-ch-ua': '"Chromium";v="112", "Google Chrome";v="112", "Not:A-Brand";v="99"',
        'sec-ch-ua-mobile': '?0',
        'sec-ch-ua-platform': '"Windows"',
        'sec-fetch-dest': 'empty',
        'sec-fetch-mode': 'cors',
        'sec-fetch-site': 'same-origin',
        'user-agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/112.0.0.0 Safari/537.36',
        'x-requested-with': 'XMLHttpRequest'
    }
    dataList = []
    datas = {}
    response = requests.request("GET", url, headers=headers, data=payload).text
    data = json.loads(response)['zpData']['hotCitySites']
    for item in data[1:]:
        name = item['name']
        code = item['code']
        dataList.append({"name": name, "code": code})
    # print(dataList)
    return dataList

def get_ip_1():
  # 提取代理API接口，获取1个代理IP
  api_url = "http://v2.api.juliangip.com/dynamic/getips?num=1&pt=1&result_type=text&split=1&trade_no=1324587523689164&sign=353a9894a22c17ae524f0ae6e412b44c"

  # 获取API接口返回的代理IP
  proxy_ip = requests.get(api_url).text

  # 用户名密码认证(动态代理/独享代理)
  username = "18038573677"
  password = "jti9iyhc"
  proxies = {
    "http": "http://%(user)s:%(pwd)s@%(proxy)s/" % {"user": username, "pwd": password, "proxy": proxy_ip},
    "https": "http://%(user)s:%(pwd)s@%(proxy)s/" % {"user": username, "pwd": password, "proxy": proxy_ip}
  }
  return proxy_ip
# mobie()
def good_link():
  dataList = []
  dataList_1 = []
  dataList_2 = []
  data_dict = [{'北京': 101010100, '上海': 101020100, '广州': 101280100, '深圳': 101280600, '杭州': 101210100, '天津': 101030100, '西安': 101110100, '苏州': 101190400, '武汉': 101200100, '厦门': 101230200, '长沙': 101250100, '成都': 101270100, '郑州': 101180100, '重庆': 101040100, '佛山': 101280800, '合肥': 101220100, '济南': 101120100, '青岛': 101120200, '南京': 101190100, '东莞': 101281600, '昆明': 101290100, '南昌': 101240100, '石家庄': 101090100, '宁波': 101210400, '福州': 101230100, '南通': 101190500, '无锡': 101190200, '珠海': 101280700, '南宁': 101300100, '常州': 101191100, '沈阳': 101070100, '大连': 101070200, '贵阳': 101260100, '惠州': 101280300, '太原': 101100100, '中山': 101281700, '泉州': 101230500, '温州': 101210700, '金华': 101210900, '海口': 101310100, '长春': 101060100, '徐州': 101190800, '哈尔滨': 101050100, '乌鲁木齐': 101130100, '嘉兴': 101210300, '保定': 101090200, '汕头': 101280500, '烟台': 101120500, '潍坊': 101120600, '江门': 101281100}]
  for item in data_dict[0].items():
    dataList.append({item[0]:item[1]})
    dataList_1.append({item[0]})
    dataList_2.append({item[1]})
    print(dataList_1)
  # return dataList


post = input('输入你需要获取的职位')
def get_url_link(name,code):
    # proxy = get_ip_1()
    chrome_options = webdriver.ChromeOptions()
    chrome_driver = r"./chromedriver.exe"
    # chrome_options.add_argument(f"--proxy-server=http://{str(proxy)}")
    # driver = webdriver.Chrome(executable_path=chrome_driver, chrome_options=chrome_options)
    driver = webdriver.Chrome(options=chrome_options, executable_path=chrome_driver)

    website = 'boss'
    sql_table(website,name,post)
    nationwide('boss','全国',post)
    for page in range(1,11):#打断点 人工过验证码
        driver.get(f'https://www.zhipin.com/web/geek/job?query={post}&city={code}&page={page}')
        time.sleep(15)
        if len(driver.find_elements(By.XPATH,'//div[@class="job-card-body clearfix"]/a')) == 0:
           break
        print(f'当前城市为{name}')
        try:
            url_list = driver.find_elements(By.XPATH,'//div[@class="job-card-body clearfix"]/a')
            diqu_list = driver.find_elements(By.XPATH,'//span[@class="job-area"]')
            zhi_wei_list = driver.find_elements(By.XPATH,'//span[@class="job-name"]')
            money_list = driver.find_elements(By.XPATH,'//div[@class="job-info clearfix"]')
            jishu_list = driver.find_elements(By.XPATH,'//div[@class="job-card-footer clearfix"]/ul')
            daiyu_list = driver.find_elements(By.XPATH,'//div[@class="info-desc"]')
            q_ye_list = driver.find_elements(By.XPATH,'//ul[@class="company-tag-list"]/li')
        except:
            print('匹配出现问题')
            continue
        # proxy = get_ip_1()
        try:
            for url,diqu,zhiwei,jishu,daiyu,q_ye,money in zip(url_list,diqu_list,zhi_wei_list,jishu_list,daiyu_list,q_ye_list,money_list):
                get_url = str(url.get_attribute('href')).split('html?')[0] + 'html'
                a = str(diqu.text).replace('\n','')
                b = str(zhiwei.text).replace('\n','')
                c = str(jishu.text).replace('\n','')
                d = str(daiyu.text).replace('\n','')
                e = str(q_ye.text).replace('\n','')
                f = str(money.text).replace('\n','')
                if redis_client.hget(f"boss_data", get_url) == None:
                    try:
                        db.insert_one(f'insert into boss_全国_{post} (详情链接,所在地区,职位,所需技术,福利待遇,企业信息,工资信息) values("{str(get_url)}","{str(a)}","{str(b)}","{str(c)}","{str(d)}","{str(e)}","{str(f)}");')
                        db.insert_one(f'insert into {website}_{name}_{post} (详情链接,所在地区,职位,所需技术,福利待遇,企业信息,工资信息) values("{str(get_url)}","{str(a)}","{str(b)}","{str(c)}","{str(d)}","{str(e)}","{str(f)}");')
                        print('插入数据成功')
                        redis_client.hset(f"boss_data", get_url,"1")
                    except:
                        continue
                else:
                    print('数据库存在此数据，跳过')
                    continue
            # proxy = get_ip_1()
        except:
            print(f'循环出现问题url==={url}')
            continue
    driver.close()

if __name__ == '__main__':
    # dataList = [{'name': '北京', 'code': 101010100}, {'name': '上海', 'code': 101020100}, {'name': '广州', 'code': 101280100}, {'name': '深圳', 'code': 101280600}, {'name': '杭州', 'code': 101210100}, {'name': '天津', 'code': 101030100}, {'name': '西安', 'code': 101110100}, {'name': '苏州', 'code': 101190400}, {'name': '武汉', 'code': 101200100}, {'name': '厦门', 'code': 101230200}, {'name': '长沙', 'code': 101250100}, {'name': '成都', 'code': 101270100}, {'name': '郑州', 'code': 101180100}, {'name': '重庆', 'code': 101040100}, {'name': '佛山', 'code': 101280800}, {'name': '合肥', 'code': 101220100}, {'name': '济南', 'code': 101120100}, {'name': '青岛', 'code': 101120200}, {'name': '南京', 'code': 101190100}, {'name': '东莞', 'code': 101281600}, {'name': '昆明', 'code': 101290100}, {'name': '南昌', 'code': 101240100}, {'name': '石家庄', 'code': 101090100}, {'name': '宁波', 'code': 101210400}, {'name': '福州', 'code': 101230100}]
    dataList = [{'name': '上海', 'code': 101020100}, {'name': '深圳', 'code': 101280600}, {'name': '杭州', 'code': 101210100}, {'name': '天津', 'code': 101030100}]
    for data in dataList:
        # print(data['name'])
        # print(data['code'])
        get_url_link(data['name'],data['code'])


